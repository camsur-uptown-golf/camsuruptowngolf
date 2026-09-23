import json
import os
import re
from collections import Counter, defaultdict
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[2]
WORKBOOK_PATH = ROOT / "test" / "outputs" / "CamSur_Uptown_Golf_QA_TestCases.xlsx"
REPORT_PATH = ROOT / "test" / "automation" / "test-results" / "results.json"
RUN_DATE = "23 September 2026"


def iter_specs(suite, inherited_file=None):
    file_name = suite.get("file") or inherited_file
    for spec in suite.get("specs", []):
        yield file_name, spec
    for child in suite.get("suites", []):
        yield from iter_specs(child, file_name)


def plain_error(test):
    results = test.get("results", [])
    if not results:
        return ""
    latest = results[-1]
    message = (latest.get("error") or {}).get("message", "")
    if not message:
        errors = latest.get("errors") or []
        if errors:
            message = errors[0].get("message", "")
    return re.sub(r"\x1b\[[0-9;]*m", "", message)


def workbook_status(test):
    reported = test.get("status")
    results = test.get("results", [])
    actual = results[-1].get("status") if results else None
    if reported == "skipped" or actual == "skipped":
        return "Skipped"
    if reported == "expected" and actual == "passed":
        return "Pass"
    return "Fail"


def result_note(file_name, title, status, error):
    if status == "Pass":
        return f"Passed in Opera GX on {RUN_DATE}."
    if status == "Skipped":
        return "Skipped: no skip-to-content link was found."
    text = f"{title}\n{error}".lower()
    if "horizontal overflow" in text:
        return "Failed in Opera GX: horizontal overflow exceeded 2 px at 320×568."
    if "tohavescreenshot" in text or "screenshot" in text:
        return "Failed visual regression: the Opera GX screenshot differed from the approved baseline."
    if "contrast" in text:
        return "Failed WCAG 2.1 AA color-contrast check; see the Playwright HTML report for affected elements."
    if "reduced motion" in text:
        return "Failed reduced-motion check because the expected animated element could not be located."
    if "locator" in text or "tobevisible" in text:
        if "mobile" in title.lower():
            return "Failed navigation check: the expected mobile drawer or submenu remained hidden."
        return "Failed navigation check: the expected menu element was not uniquely visible; selector refinement is needed."
    return f"Failed in Opera GX. See the Playwright HTML report for details ({file_name})."


def read_results():
    report = json.loads(REPORT_PATH.read_text(encoding="utf-8"))
    specs_by_file = defaultdict(list)
    for top_suite in report["suites"]:
        for file_name, spec in iter_specs(top_suite):
            specs_by_file[file_name].append(spec)

    expected_counts = {
        "task01-content-spelling.spec.ts": 42,
        "task03-responsive-layouts.spec.ts": 56,
        "task04-header-mega-nav.spec.ts": 30,
        "task10-accessibility.spec.ts": 18,
    }
    actual_counts = {name: len(specs_by_file[name]) for name in expected_counts}
    if actual_counts != expected_counts:
        raise RuntimeError(f"Unexpected Playwright spec counts: {actual_counts}")

    mapped = {}
    sequential_ranges = {
        "task10-accessibility.spec.ts": range(289, 307),
    }
    navigation_ids = {
        "Mega-menu opens — GOLF": "TC-163",
        "Mega-menu contents correct — GOLF": "TC-164",
        "Mega-menu closes on Escape — GOLF": "TC-165",
        "Mega-menu opens — CLUBHOUSE": "TC-166",
        "Mega-menu contents correct — CLUBHOUSE": "TC-167",
        "Mega-menu closes on Escape — CLUBHOUSE": "TC-168",
        "Mega-menu opens — PACKAGES": "TC-169",
        "Mega-menu contents correct — PACKAGES": "TC-170",
        "Mega-menu closes on Escape — PACKAGES": "TC-171",
        "Mega-menu opens — ACCOMMODATIONS": "TC-172",
        "Mega-menu contents correct — ACCOMMODATIONS": "TC-173",
        "Mega-menu closes on Escape — ACCOMMODATIONS": "TC-174",
        "Mega-menu opens — EXPERIENCES": "TC-175",
        "Mega-menu contents correct — EXPERIENCES": "TC-176",
        "Mega-menu closes on Escape — EXPERIENCES": "TC-177",
        "Mega-menu opens — EVENTS": "TC-178",
        "Mega-menu contents correct — EVENTS": "TC-179",
        "Mega-menu closes on Escape — EVENTS": "TC-180",
        "Mobile — hamburger opens the drawer": "TC-181",
        "Mobile — hamburger closes the drawer": "TC-182",
        "Mobile — submenu expand/collapse": "TC-183",
        "Mobile — navigating a link auto-closes the drawer": "TC-184",
        "Mobile — GOLF holes list is reachable": "TC-185",
        "Mobile — phone link uses tel: scheme": "TC-186",
        "Header is sticky on scroll": "TC-187",
        "Logo returns to Home from an interior page": "TC-188",
        "Plan-your-visit CTA navigates correctly": "TC-189",
        "Active state reflects current section": "TC-190",
        "Scroll-lock while a mega-menu is open": "TC-191",
        "Keyboard can open a mega-menu (Enter/Space)": "TC-192",
    }
    for file_name, specs in specs_by_file.items():
        if file_name == "task04-header-mega-nav.spec.ts":
            ids = [navigation_ids[spec["title"]] for spec in specs]
        elif file_name in sequential_ranges:
            ids = [f"TC-{number:03d}" for number in sequential_ranges[file_name]]
        else:
            ids = []
            for spec in specs:
                match = re.search(r"\bTC-(\d{3})\b", spec["title"])
                if not match:
                    raise RuntimeError(f"No test case ID in {file_name}: {spec['title']}")
                ids.append(f"TC-{int(match.group(1)):03d}")

        for test_case_id, spec in zip(ids, specs, strict=True):
            test = spec["tests"][0]
            status = workbook_status(test)
            mapped[test_case_id] = {
                "status": status,
                "note": result_note(file_name, spec["title"], status, plain_error(test)),
            }

    counts = Counter(item["status"] for item in mapped.values())
    if len(mapped) != 146 or counts != Counter({"Pass": 121, "Fail": 24, "Skipped": 1}):
        raise RuntimeError(f"Unexpected mapped results: total={len(mapped)}, counts={dict(counts)}")
    return mapped, counts


def resize_conditional_formatting(sheet, last_row):
    original = []
    for conditional_format, rules in sheet.conditional_formatting._cf_rules.items():
        original.append((str(conditional_format.sqref), [copy(rule) for rule in rules]))
    sheet.conditional_formatting._cf_rules.clear()
    for old_range, rules in original:
        new_range = f"J3:J{last_row}" if old_range.startswith("J3:J") else f"D3:D{last_row}"
        for rule in rules:
            sheet.conditional_formatting.add(new_range, rule)

    sheet.conditional_formatting.add(
        f"J3:J{last_row}",
        CellIsRule(
            operator="equal",
            formula=['"Skipped"'],
            fill=PatternFill("solid", fgColor="D9E1F2"),
            font=Font(color="203864", bold=True),
        ),
    )


def update_workbook(results, counts):
    workbook = load_workbook(WORKBOOK_PATH)
    cases = workbook["Test Cases"]

    # Remove every task except T01, T03, T04, and T10. Delete bottom-up so
    # the retained row styles, heights, and formulas shift with their cells.
    if cases.max_row > 148:
        cases.delete_rows(309, 86)   # T11-T13
        cases.delete_rows(195, 96)   # T05-T09
        cases.delete_rows(45, 64)    # T02

    if cases.max_row != 148:
        raise RuntimeError(f"Expected 148 worksheet rows after filtering, got {cases.max_row}")

    retained_ids = []
    module_counts = Counter()
    for row_number in range(3, cases.max_row + 1):
        test_case_id = cases.cell(row_number, 1).value
        module = str(cases.cell(row_number, 2).value or "")
        retained_ids.append(test_case_id)
        module_counts[module[:3]] += 1
        if test_case_id not in results:
            raise RuntimeError(f"No Playwright result mapped for {test_case_id}")
        cases.cell(row_number, 10).value = results[test_case_id]["status"]
        cases.cell(row_number, 11).value = results[test_case_id]["note"]
        cases.cell(row_number, 11).alignment = copy(cases.cell(row_number, 11).alignment)
        cases.cell(row_number, 11).alignment = cases.cell(row_number, 11).alignment.copy(wrap_text=True, vertical="top")

    expected_modules = Counter({"T01": 42, "T03": 56, "T04": 30, "T10": 18})
    if module_counts != expected_modules or len(retained_ids) != 146 or len(set(retained_ids)) != 146:
        raise RuntimeError(f"Retained workbook cases do not match code: {module_counts}")

    cases.auto_filter.ref = "A2:K148"
    for validation in cases.data_validations.dataValidation:
        current = str(validation.sqref)
        if current.startswith("J3:J"):
            validation.sqref = "J3:J148"
            validation.formula1 = '"Not Run,Pass,Fail,Skipped,Blocked,N/A"'
        elif current.startswith("D3:D"):
            validation.sqref = "D3:D148"
    resize_conditional_formatting(cases, 148)

    dashboard = workbook["Dashboard"]
    dashboard["B4"] = "Golf / Main Website  ·  Focused QA  ·  Environment: Opera GX / http://localhost:3000"
    dashboard["B6"] = 146
    dashboard["C6"] = 4
    dashboard["D6"] = 42
    dashboard["E6"] = 1
    dashboard["E7"] = "BROWSER"

    task_rows = [
        (1, "Review Website Copy and Spelling", "Content", "Medium", "T01*"),
        (2, "Review Accessibility and Keyboard Navigation", "Accessibility", "Medium", "T10*"),
        (3, "Test Desktop and Mobile Layouts", "UI / Responsive", "High", "T03*"),
        (4, "Test Header and Mega Navigation", "Functional", "High", "T04*"),
    ]
    for row_number, values in zip(range(11, 15), task_rows, strict=True):
        number, name, test_type, priority, pattern = values
        dashboard.cell(row_number, 2).value = number
        dashboard.cell(row_number, 3).value = name
        dashboard.cell(row_number, 4).value = test_type
        dashboard.cell(row_number, 5).value = priority
        dashboard.cell(row_number, 6).value = f'=COUNTIF(\'Test Cases\'!$B:$B,"{pattern}")'
    for row_number in range(15, 24):
        for column in range(2, 7):
            dashboard.cell(row_number, column).value = None
        dashboard.row_dimensions[row_number].hidden = True
    dashboard["F24"] = "=SUM(F11:F14)"

    dashboard["B27"] = "Not Run"
    dashboard["C27"] = "Pass"
    dashboard["D27"] = "Fail"
    dashboard["E27"] = "Skipped"
    dashboard["F27"] = "% Complete"
    dashboard["B28"] = '=COUNTIF(\'Test Cases\'!$J:$J,"Not Run")'
    dashboard["C28"] = '=COUNTIF(\'Test Cases\'!$J:$J,"Pass")'
    dashboard["D28"] = '=COUNTIF(\'Test Cases\'!$J:$J,"Fail")'
    dashboard["E28"] = '=COUNTIF(\'Test Cases\'!$J:$J,"Skipped")'
    dashboard["F28"] = "=IFERROR((C28+D28+E28)/146,0)"
    dashboard["B30"] = (
        f"Notes: Focused Opera GX run on {RUN_DATE}. The workbook contains 146 automated test cases "
        f"across the four selected QA areas. Results: {counts['Pass']} passed, {counts['Fail']} failed, "
        f"{counts['Skipped']} skipped. A separate CSpell scan checked 15 source files and reported "
        "549 unknown words in 13 files; most are Filipino or project-specific terms and need dictionary review."
    )

    legend = workbook["Legend & Scope"]
    legend["C5"] = "The focused automated suite. Filter by Module, Priority, Type or Status via the header dropdowns."
    legend["C6"] = "Status is populated from the Opera GX Playwright run: Pass / Fail / Skipped."
    legend["C9"] = (
        "146 automated test cases covering only copy and spelling, accessibility and keyboard navigation, "
        "desktop/mobile layouts, and header/mega navigation."
    )
    legend["C10"] = (
        "Opera GX against http://localhost:3000 on 23 September 2026: "
        "121 passed, 24 failed, 1 skipped."
    )

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    temporary_path = WORKBOOK_PATH.with_suffix(".tmp.xlsx")
    workbook.save(temporary_path)
    workbook.close()
    os.replace(temporary_path, WORKBOOK_PATH)


if __name__ == "__main__":
    mapped_results, status_counts = read_results()
    update_workbook(mapped_results, status_counts)
    print(f"Updated {WORKBOOK_PATH}")
    print(f"Test cases: {len(mapped_results)}")
    print(f"Statuses: {dict(status_counts)}")
