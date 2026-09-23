from copy import copy
from pathlib import Path

from openpyxl import load_workbook


book_path = Path("../../test/outputs/CamSur_Uptown_Golf_QA_TestCases.xlsx")
workbook = load_workbook(book_path, data_only=False)

for sheet in workbook.worksheets:
    print(f"SHEET {sheet.title!r}: {sheet.max_row} rows x {sheet.max_column} cols")
    print("  freeze:", sheet.freeze_panes, "merged:", list(sheet.merged_cells.ranges))
    print("  tables:", {name: table.ref for name, table in sheet.tables.items()})

print("\nDASHBOARD A1:G32")
dashboard = workbook["Dashboard"]
for row in dashboard.iter_rows(min_row=1, max_row=32, min_col=1, max_col=7):
    values = [cell.value for cell in row]
    if any(value is not None for value in values):
        print(row[0].row, values)

test_cases = workbook["Test Cases"]
print("\nTEST CASE HEADERS / SAMPLES")
for row_num in list(range(1, 10)) + list(range(10, 18)) + [108, 109, 164, 165, 280, 281, 371, 383, 394]:
    print(row_num, [test_cases.cell(row_num, col).value for col in range(1, 12)])

print("\nCOLUMN WIDTHS")
for key, dim in test_cases.column_dimensions.items():
    print(key, dim.width)

print("\nCONDITIONAL FORMATTING")
for item in test_cases.conditional_formatting:
    print(item)

print("\nDATA VALIDATIONS")
for validation in test_cases.data_validations.dataValidation:
    print(validation.sqref, validation.type, validation.formula1)

print("\nFORMULAS")
for sheet in workbook.worksheets:
    formula_cells = []
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formula_cells.append((cell.coordinate, cell.value))
    print(sheet.title, formula_cells[:100], "count=", len(formula_cells))

print("\nEXCEL")
for candidate in (
    Path(r"C:\Program Files\Microsoft Office\root\Office16\EXCEL.EXE"),
    Path(r"C:\Program Files (x86)\Microsoft Office\root\Office16\EXCEL.EXE"),
):
    print(candidate, candidate.exists())

print("\nTASK RANGES")
task_rows = {}
for row_num in range(3, test_cases.max_row + 1):
    module = str(test_cases.cell(row_num, 2).value or "")
    key = module[:3]
    task_rows.setdefault(key, []).append(row_num)
for key, rows in sorted(task_rows.items()):
    print(
        key,
        len(rows),
        rows[0],
        test_cases.cell(rows[0], 1).value,
        rows[-1],
        test_cases.cell(rows[-1], 1).value,
    )

for row_num in range(275, 305):
    print(row_num, [test_cases.cell(row_num, col).value for col in (1, 2, 3, 10)])
